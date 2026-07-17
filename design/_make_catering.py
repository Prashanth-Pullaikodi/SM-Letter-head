"""Format the 3-day catering plan into a branded SandalMist DOCX."""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LOGO = "/home/user/SM-Letter-head/design/sandalmist_logo.png"
OUT = "/home/user/SM-Letter-head/design/SandalMist_Catering_Plan.docx"

SLATE = RGBColor(0x39, 0x43, 0x4b)
COPPER = RGBColor(0xb0, 0x6a, 0x3a)
GOLD = RGBColor(0xc9, 0xa2, 0x4b)
INK = RGBColor(0x22, 0x30, 0x3c)
TEXT = RGBColor(0x41, 0x3d, 0x37)
MUTED = RGBColor(0x8f, 0x89, 0x7f)

PLAN = """First Day - 150 pax
 Evening Tea Coffee with Pazhampori
Dinner : Chappathi, Kerala Porotta, Egg Curry, Veg Kuruma
Second Day - 150 pax
Bed Coffee ( 70 pax),
Breakfast : Cut Fruits, Dosa, Upma, Sambar, Chutney, Conflex with Milk, Tea Coffee
Morning Tea Coffee with Ela Ada
Lunch : Mini Meals with fish curry
Evening Tea Coffee with Onion Pakora
Dinner : Ghee Rice, Chicken Varutharachathu, Mix Veg Curry
Third Day - 150 pax
Bed Coffee ( 70 pax),
Breakfast : Cut Fruits, Idli, Puttu, Kadala Curry, Sambar, Chutney, Conflex with Milk, Tea, Coffee
Morning Tea Coffee Vada & Chutney
Lunch : Chicken Biriyani, Veg Biriyani
Evening Kappa with Chutney, Tea Coffee"""


def pbottom(p, color="C9A24B", size=14):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bo = OxmlElement('w:bottom')
    bo.set(qn('w:val'), 'single'); bo.set(qn('w:sz'), str(size)); bo.set(qn('w:space'), '6'); bo.set(qn('w:color'), color)
    b.append(bo); pPr.append(b)

def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), hexc); tcPr.append(sh)

def no_borders(t):
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + e); el.set(qn('w:val'), 'none'); b.append(el)
    tblPr.append(b)

def run(p, text, size=10.5, color=TEXT, bold=False, italic=False, font="Arial"):
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color
    r.font.bold = bold; r.font.italic = italic; r.font.name = font; return r


doc = Document()
doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(10.5)
sec = doc.sections[0]
sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.6); sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)

# ---- Header: logo left, title right ----
h = doc.add_table(rows=1, cols=2); no_borders(h)
h.columns[0].width = Cm(4.6); h.columns[1].width = Cm(12.4)
if os.path.exists(LOGO):
    h.rows[0].cells[0].paragraphs[0].add_run().add_picture(LOGO, width=Cm(4.2))
tc = h.rows[0].cells[1].paragraphs[0]; tc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(tc, "Catering Plan", size=18, color=SLATE, bold=True, font="Georgia")
p2 = h.rows[0].cells[1].add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p2, "SandalMist Resort & Spa  ·  The Hilltop Habitat", size=8.5, color=MUTED)

rule = doc.add_paragraph(); rule.paragraph_format.space_before = Pt(4); rule.paragraph_format.space_after = Pt(8)
pbottom(rule, "C9A24B", 18)

DAY_RE = re.compile(r'(day)\s*-\s*(\d+)\s*pax', re.I)

for raw in PLAN.split("\n"):
    line = raw.strip()
    if not line:
        continue
    if DAY_RE.search(line):
        # Day heading: shaded copper bar
        t = doc.add_table(rows=1, cols=1); no_borders(t)
        cell = t.rows[0].cells[0]; shade(cell, "F3EAE1")
        pp = cell.paragraphs[0]; pp.paragraph_format.space_before = Pt(6); pp.paragraph_format.space_after = Pt(3)
        pp.paragraph_format.left_indent = Pt(2)
        # split "First Day - 150 pax" -> bold day, muted pax
        m = re.match(r'(.*day)\s*-\s*(.*)', line, re.I)
        if m:
            run(pp, m.group(1).strip().upper() + "  ", size=12, color=COPPER, bold=True, font="Georgia")
            run(pp, "— " + m.group(2).strip(), size=10, color=MUTED, bold=False)
        else:
            run(pp, line.upper(), size=12, color=COPPER, bold=True, font="Georgia")
    else:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Pt(10)
        if " : " in line:
            label, rest = line.split(" : ", 1)
            run(p, label.strip() + " : ", size=10.5, color=INK, bold=True)
            run(p, rest.strip(), size=10.5, color=TEXT)
        else:
            run(p, line, size=10.5, color=TEXT)

# footer note
fp = doc.add_paragraph(); fp.paragraph_format.space_before = Pt(18); pbottom(fp, "C9A24B", 8)
fn = doc.add_paragraph(); fn.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fn, "SandalMist Resort & Spa  ·  www.sandalmistresort.com  ·  +91 97784 34442", size=8, color=MUTED, italic=True)

doc.save(OUT)
print("saved", OUT)
