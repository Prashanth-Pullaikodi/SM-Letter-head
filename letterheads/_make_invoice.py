"""Generate a SandalMist Guest Invoice .docx template with placeholders."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/SM-Letter-head/letterheads"
LOGO = "/tmp/claude-0/-home-user-SM-Letter-head/00d8b480-5d9a-50a7-9059-f6ed2856c44a/scratchpad/logo_placeholder.png"
os.makedirs(OUT, exist_ok=True)

GREEN = RGBColor(0x1F, 0x4D, 0x46)
GOLD = RGBColor(0xC9, 0xA2, 0x4B)
CHARCOAL = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x70, 0x70, 0x70)

ADDRESS = "12/374, Sankar Hills, Munnad, Kasaragod, Kerala, 671541, India"
PHONE = "+91 97784 34442"
EMAIL = "sandalmistresortandspa@gmail.com  |  info@sandalmistresort.com"
GSTIN = "32AARFB1365N1Z6"


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_color)
    tcPr.append(sh)

def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement('w:' + e); el.set(qn('w:val'), 'none'); b.append(el)
    tblPr.append(b)

def pbottom(p, color="C9A24B", size=14):
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),str(size))
    bottom.set(qn('w:space'),'6'); bottom.set(qn('w:color'),color); pbdr.append(bottom); pPr.append(pbdr)

def run(p, text, size=10, color=CHARCOAL, bold=False, italic=False, font="Calibri"):
    r = p.add_run(text); r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; r.font.italic=italic; r.font.name=font
    return r

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10)
sec = doc.sections[0]
sec.top_margin = Cm(1.4); sec.bottom_margin = Cm(1.6); sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)

# ---------- Header: logo + company block ----------
htbl = doc.add_table(rows=1, cols=2); no_borders(htbl)
htbl.columns[0].width = Cm(3.2); htbl.columns[1].width = Cm(14.2)
lc, tc = htbl.rows[0].cells
if os.path.exists(LOGO):
    lc.paragraphs[0].add_run().add_picture(LOGO, width=Cm(2.8))
p = tc.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after = Pt(0)
run(p, "SandalMist Resort & Spa", size=20, color=GREEN, bold=True, font="Georgia")
p2 = tc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p2.paragraph_format.space_after = Pt(0)
run(p2, ADDRESS, size=8, color=GREY)
p3 = tc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p3.paragraph_format.space_after = Pt(0)
run(p3, "Phone: " + PHONE + "    " + EMAIL, size=8, color=GREY)
p4 = tc.add_paragraph(); p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p4, "GSTIN: " + GSTIN, size=8, color=GOLD, bold=True)

rule = doc.add_paragraph(); rule.paragraph_format.space_after = Pt(6); pbottom(rule, "C9A24B", 18)

# ---------- INVOICE title + meta ----------
title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_after = Pt(8)
run(title, "INVOICE", size=18, color=GREEN, bold=True, font="Georgia")

meta = doc.add_table(rows=1, cols=2); no_borders(meta)
mp = meta.rows[0].cells[0].paragraphs[0]
run(mp, "Invoice No: ", size=10, bold=True, color=GREEN); run(mp, "{INVOICE_NO}", size=10)
mp2 = meta.rows[0].cells[1].paragraphs[0]; mp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(mp2, "Date: ", size=10, bold=True, color=GREEN); run(mp2, "{INVOICE_DATE}", size=10)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------- Guest details box ----------
g = doc.add_table(rows=3, cols=2); no_borders(g)
def gcell(cell, label, tag):
    pp = cell.paragraphs[0]
    run(pp, label + ": ", size=10, bold=True, color=CHARCOAL); run(pp, "{" + tag + "}", size=10)
gcell(g.rows[0].cells[0], "Guest Name", "GUEST_NAME")
gcell(g.rows[0].cells[1], "Meal Plan", "MEAL_PLAN")
gcell(g.rows[1].cells[0], "Check-in", "CHECK_IN")
gcell(g.rows[1].cells[1], "Check-out", "CHECK_OUT")
gcell(g.rows[2].cells[0], "No. of Rooms", "NUM_ROOMS")
gcell(g.rows[2].cells[1], "No. of Guests", "NUM_GUESTS")

doc.add_paragraph().paragraph_format.space_after = Pt(4)

def section_head(text):
    t = doc.add_table(rows=1, cols=1); no_borders(t)
    c = t.rows[0].cells[0]; shade(c, "1F4D46")
    pp = c.paragraphs[0]; pp.paragraph_format.space_before = Pt(3); pp.paragraph_format.space_after = Pt(3)
    r = pp.add_run("  " + text); r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)

# ---------- Room section ----------
section_head("ROOM CHARGES")
rp = doc.add_paragraph(); rp.paragraph_format.space_before = Pt(4)
run(rp, "{ROOM_DETAILS}", size=10)
rt = doc.add_paragraph(); rt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(rt, "Room Total: Rs. ", size=10, bold=True, color=GREEN); run(rt, "{ROOM_TOTAL}", size=10, bold=True)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------- Food section ----------
section_head("FOOD CHARGES")
fp = doc.add_paragraph(); fp.paragraph_format.space_before = Pt(4)
run(fp, "{FOOD_DETAILS}", size=10)
ft = doc.add_paragraph(); ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(ft, "Food Total: Rs. ", size=10, bold=True, color=GREEN); run(ft, "{FOOD_TOTAL}", size=10, bold=True)

# ---------- Totals box ----------
doc.add_paragraph().paragraph_format.space_after = Pt(4)
tot = doc.add_table(rows=3, cols=2); no_borders(tot)
tot.alignment = WD_TABLE_ALIGNMENT.RIGHT
def totrow(row, label, tag, big=False):
    lc = row.cells[0].paragraphs[0]; lc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(lc, label, size=11 if big else 10, bold=True, color=GREEN if big else CHARCOAL)
    vc = row.cells[1].paragraphs[0]; vc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(vc, "Rs. {" + tag + "}", size=11 if big else 10, bold=True, color=GREEN if big else CHARCOAL)
totrow(tot.rows[0], "Grand Total: ", "GRAND_TOTAL")
totrow(tot.rows[1], "Advance Paid: ", "ADVANCE_PAID")
totrow(tot.rows[2], "BALANCE PAYABLE: ", "BALANCE", big=True)

# ---------- Footer ----------
foot = doc.add_paragraph(); foot.paragraph_format.space_before = Pt(20); pbottom(foot, "C9A24B", 8)
fp2 = doc.add_paragraph(); fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp2, "Thank you for staying with SandalMist Resort & Spa - The Hill Top Habitat", size=9, italic=True, color=GREY)
fp3 = doc.add_paragraph(); fp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp3, "This is a computer-generated invoice.", size=8, color=GREY)

path = os.path.join(OUT, "05_Guest_Invoice_Template.docx")
doc.save(path)
print("saved", path)
